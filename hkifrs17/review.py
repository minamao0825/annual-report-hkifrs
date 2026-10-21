"""Excel review round trip with source and check sheets."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .data import COLUMNS, REVIEW_STATUSES, normalize_records, reconciliation_checks
from .checks import ASSESSMENTS, FORMULA_BASES, METHODS, NOTE_SCOPES, PLAN_COLUMNS, default_plan, normalize_plan


AUDIT_STATUS_OPTIONS = (
    "Check source",
    "Audited statement",
    "Comparative in 2025 statement",
    "Unaudited comparative",
)


def _add_status_dropdown(sheet, column: str, options: tuple[str, ...], *, required: bool) -> None:
    letter = get_column_letter(COLUMNS.index(column) + 1)
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(options) + '"',
        allow_blank=not required,
    )
    validation.showDropDown = False
    validation.showErrorMessage = required
    validation.errorTitle = "Choose a valid status"
    validation.error = f"Select a status from the {column} dropdown."
    sheet.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}1048576")


def review_workbook_bytes(frame: pd.DataFrame, filings: list[dict], plan: pd.DataFrame | None = None) -> bytes:
    rows = normalize_records(frame, filings)
    check_plan = normalize_plan(plan, filings, set(rows["Document ID"]))
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        rows.to_excel(writer, sheet_name="Records", index=False)
        checks = reconciliation_checks(rows, filings, check_plan, verified_only=True)
        checks.to_excel(writer, sheet_name="Checks", index=False)
        check_plan.to_excel(writer, sheet_name="Check plan", index=False)
        pd.DataFrame([
            {"Instruction": "Check each original value, source unit, period and PDF page against the filing."},
            {"Instruction": "Set Review status to Verified only after source review. Excluded rows never enter charts."},
            {"Instruction": "Value (millions), Metric, Section and identity columns are recalculated on import. For percent source units, the numeric value is percentage points rather than millions."},
            {"Instruction": "A 2024 comparative does not automatically mean a separately audited 2024 filing."},
            {"Instruction": "Checks use only Verified reported cells. Missing inputs and Difference block approval."},
            {"Instruction": "Set each optional Check plan row to Applicable or Not applicable; the latter needs a reason."},
            {"Instruction": "For applicable note checks, confirm measurement method, note reference and matching scope where relevant."},
            {"Instruction": "For the SERVICE check only, set Formula basis if source expenses are shown as positive amounts."},
            {"Instruction": "A signed net movement must be independently disclosed in the same note and basis; never calculate it from opening and closing."},
        ]).to_excel(writer, sheet_name="Read me", index=False)
        for sheet in writer.book:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.column_dimensions["A"].width = 27
            sheet.column_dimensions["B"].width = 25
            if sheet.title == "Records":
                sheet.column_dimensions["P"].width = 46
                sheet.column_dimensions["S"].width = 62
                _add_status_dropdown(sheet, "Audit status", AUDIT_STATUS_OPTIONS, required=False)
                _add_status_dropdown(sheet, "Review status", REVIEW_STATUSES, required=True)
            elif sheet.title == "Check plan":
                for column, options in (("D", ASSESSMENTS), ("E", METHODS), ("F", NOTE_SCOPES), ("I", FORMULA_BASES)):
                    validation = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False)
                    validation.showErrorMessage = True
                    sheet.add_data_validation(validation)
                    validation.add(f"{column}2:{column}1048576")
                sheet.column_dimensions["G"].width = 42
                sheet.column_dimensions["H"].width = 52
                sheet.column_dimensions["I"].width = 42
    return output.getvalue()


def read_review_workbook(payload: bytes, filings: list[dict]) -> pd.DataFrame:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
    if "Records" not in workbook:
        raise ValueError("The Records sheet is missing.")
    sheet = workbook["Records"]
    records = list(sheet.values)
    if not records or list(records[0])[:len(COLUMNS)] != COLUMNS:
        raise ValueError("Records headers have changed.")
    if any(cell.data_type == "f" for row in sheet.iter_rows(min_row=2) for cell in row):
        raise ValueError("Replace formulas in Records with reviewed values before upload.")
    return normalize_records(pd.DataFrame(records[1:], columns=COLUMNS), filings)


def read_check_plan_workbook(payload: bytes, filings: list[dict], document_id: str) -> pd.DataFrame:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
    if "Check plan" not in workbook:
        return default_plan(filings, {document_id})
    sheet = workbook["Check plan"]
    records = list(sheet.values)
    if not records or list(records[0])[:len(PLAN_COLUMNS)] != PLAN_COLUMNS:
        raise ValueError("Check plan headers have changed.")
    if any(cell.data_type == "f" for row in sheet.iter_rows(min_row=2) for cell in row):
        raise ValueError("Replace formulas in Check plan with reviewed values before upload.")
    return normalize_plan(pd.DataFrame(records[1:], columns=PLAN_COLUMNS), filings, {document_id})
