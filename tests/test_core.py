from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
import requests

from hkifrs17.catalog import INSURER_CHART_GROUPS, METRICS, PEER_METRIC_GROUPS, PILOT_FILINGS
from hkifrs17.data import normalize_records, reconciliation_checks
from hkifrs17.checks import RULES, default_plan, normalize_plan
from hkifrs17.demo import demo_records
from hkifrs17.charts import peer_amounts, peer_margin
from hkifrs17.review import read_check_plan_workbook, read_review_workbook, review_workbook_bytes
from hkifrs17.monitor import DEFAULT_LINKS, LINK_COLUMNS, validate_links, workbook_bytes, scan_link
import hkifrs17.storage as storage


def test_pilot_bridges_reconcile_in_both_years():
    checks = reconciliation_checks(demo_records(), list(PILOT_FILINGS))
    assert len(checks) == len(PILOT_FILINGS) * 2 * len(RULES)
    assert checks.loc[checks["Check ID"].isin(["PRIMARY", "SERVICE", "BALANCE"]), "Result"].eq("Pass").all()
    assert checks.loc[~checks["Check ID"].isin(["PRIMARY", "SERVICE", "BALANCE"]), "Result"].eq("Missing inputs").all()


def test_periods_and_audit_flags_stay_attached_to_source():
    rows = demo_records()
    assert set(rows.loc[rows["Company ID"].eq("CTF"), "Period end"]) == {"2025-06-30", "2024-06-30"}
    assert rows.loc[rows["Company ID"].eq("SUN") & rows["Period end"].eq("2024-12-31"), "Audit status"].eq("Unaudited comparative").all()


def test_currency_conflict_is_rejected():
    rows = demo_records("FWD_2025")
    pages = rows.set_index(["Metric code", "Period end"])["PDF page"]
    assert int(pages.loc[("INS_NOTE_NET_OPEN", "2025-12-31")]) == 63
    assert int(pages.loc[("INS_NOTE_NET_OPEN", "2024-12-31")]) == 64
    rows.loc[0, "Source unit"] = "HKD thousands"
    with pytest.raises(ValueError, match="unit currency conflicts"):
        normalize_records(rows, list(PILOT_FILINGS))


def test_rollforward_requires_disclosed_movement_and_flags_difference():
    filings = list(PILOT_FILINGS)
    rows = demo_records("YF_2025")
    additions = []
    for code, amount in (("CSM_OPEN", 10_000_000), ("CSM_CLOSE", 15_000_000)):
        row = rows.iloc[0].copy()
        row["Metric code"] = code
        row["Original value"] = amount
        row["PDF page"] = 33
        row["Source label"] = code
        additions.append(row)
    rows = normalize_records(pd.concat([rows, pd.DataFrame(additions)], ignore_index=True), filings)
    plan = default_plan(filings, {"YF_2025"})
    optional = ~plan["Check ID"].isin(["PRIMARY", "SERVICE", "BALANCE"])
    plan.loc[optional, "Assessment"] = "Not applicable"
    plan.loc[optional, "Reason"] = "No comparable note disclosure."
    csm = plan["Check ID"].eq("CSM_ROLL") & plan["Period end"].eq("2025-12-31")
    plan.loc[csm, ["Assessment", "Measurement method", "Note scope", "Note reference", "Reason"]] = [
        "Applicable", "GMM", "Different or net scope", "Note 13, page 33", ""]
    checks = reconciliation_checks(rows, filings, plan)
    result = checks.loc[checks["Check ID"].eq("CSM_ROLL") & checks["Period end"].eq("2025-12-31")].iloc[0]
    assert result["Result"] == "Missing inputs"
    assert result["Missing inputs"] == "CSM_NET_MOVEMENT"

    movement = additions[0].copy()
    movement["Metric code"] = "CSM_NET_MOVEMENT"
    movement["Original value"] = 5_000_000
    movement["Source label"] = "Disclosed signed net movements"
    rows = normalize_records(pd.concat([rows, pd.DataFrame([movement])], ignore_index=True), filings)
    checks = reconciliation_checks(rows, filings, plan)
    assert checks.loc[checks["Check ID"].eq("CSM_ROLL") & checks["Period end"].eq("2025-12-31"), "Result"].iloc[0] == "Pass"
    rows.loc[rows["Metric code"].eq("CSM_CLOSE") & rows["Period end"].eq("2025-12-31"), "Original value"] = 16_000_000
    checks = reconciliation_checks(rows, filings, plan)
    assert checks.loc[checks["Check ID"].eq("CSM_ROLL") & checks["Period end"].eq("2025-12-31"), "Result"].iloc[0] == "Difference"


def test_cross_table_requires_confirmed_matching_scope_and_na_reason():
    filings = list(PILOT_FILINGS)
    rows = demo_records("YF_2025")
    note = rows.loc[rows["Metric code"].eq("INS_REVENUE") & rows["Period end"].eq("2025-12-31")].iloc[0].copy()
    note["Metric code"] = "INS_REVENUE_NOTE_TOTAL"
    note["PDF page"] = 40
    note["Source label"] = "Total insurance revenue in note"
    rows = normalize_records(pd.concat([rows, pd.DataFrame([note])], ignore_index=True), filings)
    plan = default_plan(filings, {"YF_2025"})
    target = plan["Check ID"].eq("REVENUE_NOTE") & plan["Period end"].eq("2025-12-31")
    plan.loc[target, ["Assessment", "Measurement method", "Note reference"]] = ["Applicable", "Mixed", "Note 5, page 40"]
    checks = reconciliation_checks(rows, filings, plan)
    assert checks.loc[checks["Check ID"].eq("REVENUE_NOTE") & checks["Period end"].eq("2025-12-31"), "Result"].iloc[0] == "Missing inputs"
    plan.loc[target, "Note scope"] = "Same entity and balance basis"
    checks = reconciliation_checks(rows, filings, plan)
    assert checks.loc[checks["Check ID"].eq("REVENUE_NOTE") & checks["Period end"].eq("2025-12-31"), "Result"].iloc[0] == "Pass"
    plan.loc[target, "Assessment"] = "Not applicable"
    with pytest.raises(ValueError, match="explain why"):
        normalize_plan(plan, filings, {"YF_2025"})


def test_service_bridge_can_use_positive_expense_presentation():
    filings = list(PILOT_FILINGS)
    rows = demo_records("YF_2025")
    rows.loc[rows["Metric code"].eq("INS_EXPENSE"), "Original value"] *= -1
    plan = default_plan(filings, {"YF_2025"})
    service = plan["Check ID"].eq("SERVICE")
    plan.loc[service, "Formula basis"] = "Subtract positive service expense"
    checks = reconciliation_checks(rows, filings, plan)
    assert checks.loc[checks["Check ID"].eq("SERVICE"), "Result"].eq("Pass").all()
    workbook = review_workbook_bytes(rows, filings, plan)
    imported = read_check_plan_workbook(workbook, filings, "YF_2025")
    assert imported.loc[imported["Check ID"].eq("SERVICE"), "Formula basis"].eq("Subtract positive service expense").all()


def test_fwd_disclosed_note_components_and_net_rollforwards_reconcile():
    filings = list(PILOT_FILINGS)
    rows = demo_records("FWD_2025")
    plan = default_plan(filings, {"FWD_2025"})
    for index, item in plan.iterrows():
        check_id = item["Check ID"]
        if check_id in {"PRIMARY", "SERVICE", "BALANCE"}:
            continue
        if check_id in {"LIABILITY_ROLL", "REINSURANCE_ROLL"}:
            plan.at[index, "Assessment"] = "Not applicable"
            plan.at[index, "Reason"] = "Gross movement is not a separately disclosed total in this selected note cohort."
            continue
        plan.at[index, "Assessment"] = "Applicable"
        plan.at[index, "Measurement method"] = "Mixed"
        plan.at[index, "Note scope"] = "Same entity and balance basis" if check_id in {
            "REVENUE_NOTE", "LIABILITY_NOTE", "REINSURANCE_NOTE"} else "Different or net scope"
        plan.at[index, "Note reference"] = "Notes 4 and 13, PDF pp.51–74"
        if check_id == "LIABILITY_NOTE":
            plan.at[index, "Formula basis"] = "Reverse note sign for statement comparison"
    checks = reconciliation_checks(rows, filings, plan)
    assert checks["Result"].isin(["Pass", "Not applicable"]).all(), checks.to_string(index=False)
    assert checks["Result"].eq("Pass").sum() == 30
    assert checks["Result"].eq("Not applicable").sum() == 4


def test_new_metric_charts_cover_dictionary_and_fwd_text_note():
    from hkifrs17.charts import component_bars, note_movement_bridge, peer_ratio
    from hkifrs17.extraction_plan import metric_codes_for
    from hkifrs17.pdf import extract_text_page
    from hkifrs17.storage import pdf_path

    chart_codes = [code for _, group in INSURER_CHART_GROUPS for code in group]
    rate_codes = {metric.code for metric in METRICS if metric.kind == "rate"}
    assert len(chart_codes) == len(set(chart_codes))
    assert set(chart_codes) == {metric.code for metric in METRICS} - rate_codes
    assert {code for _, group in PEER_METRIC_GROUPS for code in group}.issubset(set(chart_codes))
    filing = PILOT_FILINGS[1]
    note = extract_text_page(pdf_path(filing), filing, 51, list(PILOT_FILINGS),
                             metric_codes_for(["Insurance revenue note"]))
    assert note["Metric code"].nunique() == 7
    assert note.loc[note["Metric code"].eq("INS_REVENUE_NOTE_TOTAL") & note["Period end"].eq("2025-12-31"),
                    "Original value"].iloc[0] == 875076
    fwd = demo_records("FWD_2025")
    assert component_bars(fwd, "2025-12-31", ["REV_CSM_SERVICE", "REV_RA_RELEASE"]) is not None
    assert note_movement_bridge(fwd, "2025-12-31", (
        "INS_NOTE_NET_OPEN", "INS_NOTE_PL_OCI", "INS_NOTE_CASHFLOWS", "INS_NOTE_NET_CLOSE")) is not None
    assert peer_ratio(demo_records(), "2025-12-31", "INS_LIABS", "TOTAL_ASSETS", "Share (%)") is not None


def test_peer_amounts_use_matching_year_end_rate():
    rows = demo_records()
    cohort = rows[rows["Period end"].eq("2025-12-31")]
    assert peer_margin(cohort, "2025-12-31") is not None
    chart, excluded = peer_amounts(cohort, "2025-12-31", "INS_REVENUE")
    assert chart is not None
    assert excluded == []


def test_review_and_publish_gate(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA", tmp_path)
    monkeypatch.setattr(storage, "UPLOADS_FILE", tmp_path / "uploaded_filings.json")
    monkeypatch.setattr(storage, "PUBLISHED_FILE", tmp_path / "published_records.xlsx")
    draft = demo_records("YF_2025")
    storage.save_draft("YF_2025", draft)
    with pytest.raises(ValueError, match="Verify at least one"):
        storage.approve_draft("YF_2025")
    draft["Review status"] = "Verified"
    workbook = review_workbook_bytes(draft, list(PILOT_FILINGS))
    reviewed = read_review_workbook(workbook, list(PILOT_FILINGS))
    storage.save_draft("YF_2025", reviewed)
    with pytest.raises(ValueError, match="Missing inputs"):
        storage.approve_draft("YF_2025")
    plan = storage.load_check_plan("YF_2025")
    optional = ~plan["Check ID"].isin(["PRIMARY", "SERVICE", "BALANCE"])
    plan.loc[optional, "Assessment"] = "Not applicable"
    plan.loc[optional, "Reason"] = "No matching disclosed note fields in the reviewed pilot sample."
    storage.save_check_plan("YF_2025", plan)
    assert storage.approve_draft("YF_2025") == 26
    assert storage.publish_selected(["YF_2025"]) == 26
    assert len(storage.load_published()) == 26
    assert storage.integrate_published() == 26
    assert storage.integration_status()["current"]
    assert len(storage.load_integrated()) == 26
    reviewed.loc[reviewed["Metric code"].eq("OPERATING_CF"), "Review status"] = "Excluded"
    storage.save_draft("YF_2025", reviewed)
    with pytest.raises(ValueError, match="Missing inputs"):
        storage.approve_draft("YF_2025")
    assert len(storage.load_published()) == 26
    assert storage.integration_status()["current"]


def test_review_workbook_has_audit_and_review_status_dropdowns():
    from io import BytesIO
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(review_workbook_bytes(demo_records("YF_2025"), list(PILOT_FILINGS))))
    records = workbook["Records"]
    validations = {str(item.sqref): item for item in records.data_validations.dataValidation}
    assert set(validations) == {"Q2:Q1048576", "R2:R1048576"}
    assert validations["Q2:Q1048576"].formula1 == '"Check source,Audited statement,Comparative in 2025 statement,Unaudited comparative"'
    assert validations["Q2:Q1048576"].showDropDown is False
    assert validations["R2:R1048576"].formula1 == '"Pending,Verified,Excluded"'
    assert validations["R2:R1048576"].showErrorMessage is True
    assert workbook.sheetnames == ["Records", "Checks", "Check plan", "Read me"]


def test_legacy_review_workbook_import_requires_fresh_check_assessment():
    from io import BytesIO
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(review_workbook_bytes(demo_records("YF_2025"), list(PILOT_FILINGS))))
    del workbook["Check plan"]
    output = BytesIO()
    workbook.save(output)
    plan = read_check_plan_workbook(output.getvalue(), list(PILOT_FILINGS), "YF_2025")
    assert plan.loc[plan["Check ID"].eq("CSM_ROLL"), "Assessment"].eq("To assess").all()


def test_company_list_excel_round_trip_and_private_links_rejected():
    frame = validate_links(pd.DataFrame(DEFAULT_LINKS, columns=LINK_COLUMNS))
    imported = pd.read_excel(__import__("io").BytesIO(workbook_bytes(frame)), sheet_name="Insurer links")
    assert validate_links(imported)["Insurer"].tolist() == frame["Insurer"].tolist()
    assert validate_links(imported)["Peer group"].tolist() == frame["Peer group"].tolist()
    assert validate_links(imported.drop(columns="Peer group"))["Peer group"].eq("Life insurer").all()
    frame.loc[0, "Report page URL"] = "https://127.0.0.1/internal"
    with pytest.raises(ValueError, match="Private network"):
        validate_links(frame)


def test_public_page_scan_is_a_manual_confirmation_lead(monkeypatch):
    class Response:
        status_code = 200
        url = "https://example.org/financials"
        headers = {"Content-Type": "text/html"}

        def raise_for_status(self):
            pass

        def iter_content(self, _):
            yield b'<html><body>2025 Audited Financial Statements <a href="/2025-report.pdf">2025 report</a></body></html>'

        def close(self):
            pass

    monkeypatch.setattr("hkifrs17.monitor._fetch", lambda _: Response())
    result = scan_link({"Insurer": "Test", "Report page URL": "https://example.org/financials",
                        "Report terms": "audited financial statements"}, 2025)
    assert result["Status"] == "Potential filing — confirm PDF"
    assert result["Report link"] == "https://example.org/2025-report.pdf"


def test_kimi_extraction_uses_fast_mode_and_configurable_timeout(monkeypatch):
    import hkifrs17.pdf as pdf

    seen = []

    class Response:
        def raise_for_status(self):
            pass

        def json(self):
            return {"choices": [{"message": {"content": '{"records": []}'}}]}

    def fake_post(url, **kwargs):
        seen.append((url, kwargs))
        return Response()

    monkeypatch.setattr(pdf, "page_png", lambda *args, **kwargs: b"test image")
    monkeypatch.setattr(pdf.requests, "post", fake_post)
    pdf.extract_model_page(Path("unused.pdf"), PILOT_FILINGS[0], 27, list(PILOT_FILINGS),
                           base_url="https://api.moonshot.cn/v1", model="kimi-k2.6", api_key="test-key",
                           read_timeout=300)
    assert seen[0][0] == "https://api.moonshot.cn/v1/chat/completions"
    assert seen[0][1]["json"]["thinking"] == {"type": "disabled"}
    assert seen[0][1]["timeout"] == (10, 300)
    pdf.extract_model_page(Path("unused.pdf"), PILOT_FILINGS[0], 27, list(PILOT_FILINGS),
                           base_url="https://api.example.org/v1", model="another-vision-model", api_key="test-key")
    assert "thinking" not in seen[1][1]["json"]


def test_kimi_timeout_does_not_retry_or_change_draft(monkeypatch):
    import hkifrs17.pdf as pdf

    calls = []

    def slow_post(*args, **kwargs):
        calls.append(kwargs)
        raise requests.exceptions.ReadTimeout("read timed out")

    monkeypatch.setattr(pdf, "page_png", lambda *args, **kwargs: b"test image")
    monkeypatch.setattr(pdf.requests, "post", slow_post)
    with pytest.raises(TimeoutError, match="existing draft was not changed"):
        pdf.extract_model_page(Path("unused.pdf"), PILOT_FILINGS[0], 27, list(PILOT_FILINGS),
                               base_url="https://api.moonshot.cn/v1", model="kimi-k2.6", api_key="test-key")
    assert len(calls) == 1


def test_batch_plan_groups_pages_and_rejects_invalid_selection():
    from hkifrs17.extraction_plan import group_pages, metric_codes_for, suggested_plan, validate_plan

    plan = suggested_plan(PILOT_FILINGS[0], 147)
    assert plan["Extract"].tolist()[:3] == [True, True, True]
    assert not any(plan["Extract"].tolist()[3:])
    assert plan.loc[plan["Table"].eq("Equity movement"), "Physical PDF page"].tolist() == [31]
    fwd_plan = suggested_plan(PILOT_FILINGS[1], 147)
    assert fwd_plan["Extract"].tolist()[:4] == [True, True, True, False]
    assert fwd_plan.loc[fwd_plan["Table"].eq("Insurance revenue note"), "Physical PDF page"].tolist() == [51]
    plan.loc[len(plan)] = [True, 27, "Financial position"]
    selections = validate_plan(plan, 147)
    assert group_pages(selections)[27] == ["Profit or loss", "Financial position"]
    assert "OPERATING_CF" not in metric_codes_for(group_pages(selections)[27])
    plan.loc[0, "Physical PDF page"] = 148
    with pytest.raises(ValueError, match="page must be an integer"):
        validate_plan(plan, 147)


def test_batch_merge_preserves_reviewed_cells_and_refreshes_pending():
    from hkifrs17.extraction_plan import merge_extracted

    current = demo_records("YF_2025").iloc[:2].copy()
    current.loc[current.index[0], "Review status"] = "Verified"
    incoming = current.copy()
    incoming["Original value"] = [111, 222]
    incoming["Review status"] = "Pending"
    merged, counts = merge_extracted(current, incoming)
    assert counts == {"added": 0, "replaced_pending": 1, "preserved": 1}
    assert merged.loc[merged["Metric code"].eq(current.iloc[0]["Metric code"])
                      & merged["Period end"].eq(current.iloc[0]["Period end"]), "Original value"].iloc[0] == current.iloc[0]["Original value"]


def test_confirmed_batch_continues_after_table_failure_and_reports_progress():
    from hkifrs17.extraction_plan import extract_confirmed_batch

    source = demo_records("YF_2025")
    draft = [source.iloc[:0].copy()]
    calls = []
    progress = []

    def extractor(page, table):
        calls.append((page, table))
        if table == "Financial position":
            raise TimeoutError("Model timed out")
        row = 0 if table == "Profit or loss" else 2
        return source.iloc[[row]].copy()

    result = extract_confirmed_batch(
        [(27, "Profit or loss"), (29, "Financial position"), (32, "Cash flows")],
        extractor=extractor,
        load_current=lambda: draft[0],
        save_current=lambda frame: draft.__setitem__(0, frame),
        progress=lambda *item: progress.append(item),
    )
    assert calls == [(27, "Profit or loss"), (29, "Financial position"), (32, "Cash flows")]
    assert result["Status"].tolist() == ["Saved as Pending", "Failed", "Saved as Pending"]
    assert result.loc[1, "Detail"] == "Model timed out"
    assert len(draft[0]) == 2
    assert [item[0] for item in progress] == [1, 2, 3]


def test_model_batch_scope_only_returns_selected_table_metrics(monkeypatch):
    import hkifrs17.pdf as pdf

    captured = {}

    class Response:
        def raise_for_status(self):
            pass

        def json(self):
            return {"choices": [{"message": {"content": __import__("json").dumps({"records": [
                {"metric_code": "OPERATING_CF", "period_end": "2025-12-31", "original_value": 10,
                 "source_unit": "HKD", "source_label": "Operating cash flow"},
                {"metric_code": "INS_REVENUE", "period_end": "2025-12-31", "original_value": 99,
                 "source_unit": "HKD", "source_label": "Insurance revenue"},
            ]})}}]}

    def fake_post(url, **kwargs):
        captured.update(kwargs)
        return Response()

    monkeypatch.setattr(pdf, "page_png", lambda *args, **kwargs: b"page")
    monkeypatch.setattr(pdf.requests, "post", fake_post)
    rows = pdf.extract_model_page(Path("unused.pdf"), PILOT_FILINGS[0], 32, list(PILOT_FILINGS),
                                  base_url="https://api.moonshot.cn/v1", model="kimi-k2.6", api_key="test-key",
                                  metric_codes={"OPERATING_CF"}, table_names=["Cash flows"])
    assert rows["Metric code"].tolist() == ["OPERATING_CF"]
    prompt = captured["json"]["messages"][1]["content"][0]["text"]
    assert "Target table(s): Cash flows" in prompt
    assert "INS_REVENUE:" not in prompt
